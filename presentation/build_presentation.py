"""
Builds PathSense_SIH_Presentation.pptx on the OFFICIAL SIH 2026 idea-presentation template.

    python presentation/build_presentation.py  [path\\to\\SIH2026-IDEA-Presentation-Format.pptx]

This is the IDEA deck: it presents the proposed solution, its feasibility, roadmap and expected impact.
It deliberately contains no prototype measurements (tests, FPS, latency, detections, ...) and no screenshots of
the prototype; those belong to the live demo. The official format is kept: six slides including the title,
fixed headings and the official "idea details pointers" (verbatim).

Visuals: presentation/assets/hero_hero.png and hero_iso.png are ORIGINAL 3D concept renders produced by
presentation/hero/render_hero.mjs (three.js scene in presentation/hero/scene.html); the *_anchors.json files
next to them place the slide callouts on the rendered objects.
"""

import json
import os
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "presentation", "assets")
ICONS = os.path.join(ASSETS, "icons")
TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser("~"), "Downloads", "SIH2026-IDEA-Presentation-Format.pptx")
TEAM_NAME, TEAM_ID = "Ctrl Alt Elite", ""
PS_ID = "SIH26037"
PS_TITLE = "Adaptive Path Planning and Collision Avoidance for Autonomous Vehicles on Unstructured Indian Roads"

# ------------------------------------------------------------------ design tokens (white template, SIH blue)
NAVY, INK, MUTED, LINE = "1F2A44", "243044", "5B6778", "D5DDE8"
TINT, TINT2 = "F3F7FC", "EAF2FB"
BLUE, TEAL, TEAL_T = "0070C0", "0F9E8E", "E3F4F1"
AMBER, AMBER_T, RED, RED_T, GREY = "D97706", "FDF3E3", "DC2626", "FDECEC", "7B8594"
BODY = "Arial"


def rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------ primitives
prs = Presentation(TEMPLATE)


def delete_slide(idx):
    sld = prs.slides._sldIdLst[idx]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def shape_by_name(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def remove(sh):
    sh._element.getparent().remove(sh._element)


def set_text_keep_format(sh, text_):
    """replace a text frame's text but keep the first run's formatting"""
    tf = sh.text_frame
    p0 = next((p for p in tf.paragraphs if p.runs), tf.paragraphs[0])
    runs = p0.runs
    runs[0].text = text_
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs):
        if p._p is not p0._p:
            p._p.getparent().remove(p._p)


def box(s, x, y, w, h, fill=TINT, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=0.75, dash=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            etree.SubElement(sh.line._get_or_add_ln(), qn("a:prstDash")).set("val", dash)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=BODY,
         space=0, bullets=False, line_spacing=None, bullet_color=BLUE):
    """paras: str | [para]; para: str | [(text, {opts})]."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        if line_spacing:
            p.line_spacing = line_spacing
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.16)))); pPr.set("indent", str(-int(Inches(0.16))))
            bc = etree.SubElement(pPr, qn("a:buClr")); etree.SubElement(bc, qn("a:srgbClr")).set("val", bullet_color)
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def pointer(s, x, y, w, label, size=13, color=NAVY):
    """an official 'idea details pointer', kept verbatim, styled as a section label"""
    box(s, x, y + 0.05, 0.06, 0.23, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.14, y, w - 0.14, 0.32, label, size=size, bold=True, color=color)


def icon(s, name, x, y, size, color="b", circle=None, pad=0.2):
    f = os.path.join(ICONS, f"{name}_{color}.png")
    if circle:
        box(s, x, y, size, size, fill=circle, shape=MSO_SHAPE.OVAL)
        p = size * pad
        s.shapes.add_picture(f, Inches(x + p), Inches(y + p), Inches(size - 2 * p), Inches(size - 2 * p))
    else:
        s.shapes.add_picture(f, Inches(x), Inches(y), Inches(size), Inches(size))


def connector(s, x1, y1, x2, y2, color=MUTED, width=1.5, dash=None, head=True, kind=MSO_CONNECTOR.STRAIGHT):
    c = s.shapes.add_connector(kind, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", dash)
    if head:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def arrow(s, x1, y1, x2, y2, color=MUTED, width=1.5, dash=None):
    return connector(s, x1, y1, x2, y2, color, width, dash)


def line(s, x1, y1, x2, y2, color=LINE, width=1.0, dash=None):
    return connector(s, x1, y1, x2, y2, color, width, dash, head=False)


def pill(s, x, y, label, color, size=10, h=0.3, fill="FFFFFF", w=None, bold=True):
    w = w or (0.26 + 0.068 * len(label) * size / 10)
    box(s, x, y, w, h, fill=fill, line=color, radius=h / 2, lw=1.1)
    text(s, x, y, w, h, label, size=size, bold=bold, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def callout(s, ax, ay, lx, ly, label, color, size=10):
    """leader line from an anchor point (ax, ay) on the image to a pill label centred at (lx, ly)"""
    w = 0.26 + 0.068 * len(label) * size / 10
    h = 0.3
    ex = lx - w / 2 if ax < lx - w / 2 else (lx + w / 2 if ax > lx + w / 2 else ax)
    ey = ly - h / 2 if ay < ly else ly + h / 2
    line(s, ax, ay, ex, ey, color=color, width=1.0)
    box(s, ax - 0.045, ay - 0.045, 0.09, 0.09, fill=color, shape=MSO_SHAPE.OVAL)
    pill(s, lx - w / 2, ly - h / 2, label, color, size=size, w=w)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def team_badge(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.replace("\n", " ").strip().startswith("Your Team"):
            set_text_keep_format(sh, TEAM_NAME)


# ------------------------------------------------------------------ images (original renders -> slide-ready files)
def prep_images():
    """hero: crop the empty sky, fade the left/top edges to white-transparent so it bleeds into the slide.
       iso: crop to the road and save as JPEG. Returns (hero_path, hero_crop_top, iso_path, iso_crop_box)."""
    src = Image.open(os.path.join(ASSETS, "hero_hero.png")).convert("RGBA")
    W, H = src.size
    top = 0.06
    hero = src.crop((0, int(H * top), W, H))
    hero = hero.resize((2000, int(2000 * hero.height / hero.width)), Image.LANCZOS)
    w, h = hero.size
    alpha = Image.new("L", (w, h), 255)
    px = alpha.load()
    fx, fy = int(w * 0.24), int(h * 0.10)
    for x in range(w):
        ax = 1.0 if x >= fx else (x / fx) ** 1.6
        for y in range(h):
            ay = 1.0 if y >= fy else (y / fy)
            px[x, y] = int(255 * ax * ay)
    hero.putalpha(alpha)
    hero_path = os.path.join(ASSETS, "hero_slide.png")
    hero.save(hero_path, optimize=True)

    iso = Image.open(os.path.join(ASSETS, "hero_iso.png")).convert("RGB")
    W2, H2 = iso.size
    cb = (0.16, 0.10, 0.86, 0.93)
    iso = iso.crop((int(W2 * cb[0]), int(H2 * cb[1]), int(W2 * cb[2]), int(H2 * cb[3])))
    iso = iso.resize((2000, int(2000 * iso.height / iso.width)), Image.LANCZOS)
    iso_path = os.path.join(ASSETS, "hero_iso_slide.jpg")
    iso.save(iso_path, quality=90, optimize=True)
    return hero_path, top, iso_path, cb


HERO, HERO_TOP, ISO, ISO_BOX = prep_images()
A_HERO = json.load(open(os.path.join(ASSETS, "hero_hero_anchors.json")))["anchors"]
A_ISO = json.load(open(os.path.join(ASSETS, "hero_iso_anchors.json")))["anchors"]


slides = list(prs.slides)
if len(slides) == 7:          # slide 7 = "IMPORTANT INSTRUCTIONS" - the template says to delete it before upload
    delete_slide(6)
slides = list(prs.slides)
assert len(slides) == 6, len(slides)
for s in slides[1:]:
    team_badge(s)

# ================================================================== 1. TITLE PAGE (official fields)
s = slides[0]
remove(shape_by_name(s, "Picture 4"))              # template bulb artwork -> replaced by the concept render
remove(shape_by_name(s, "Freeform: Shape 26"))
# hero render: bleeds off the right and bottom edges, fades into the page on the left/top
IX, IW = 5.35, 13.333 - 5.35
img = Image.open(HERO)
IH = IW * img.height / img.width
IY = 7.5 - IH
pic = s.shapes.add_picture(HERO, Inches(IX), Inches(IY), Inches(IW), Inches(IH))
s.shapes._spTree.remove(pic._element)
s.shapes._spTree.insert(3, pic._element)           # behind the header texts / logo


def hero_xy(k):
    ax, ay = A_HERO[k]
    return IX + ax * IW, IY + (ay - HERO_TOP) / (1 - HERO_TOP) * IH


callout(s, *hero_xy("corridor"), IX + 1.55, 4.55, "Safe corridor", TEAL)
callout(s, *hero_xy("pedestrian"), IX + 3.35, 2.55, "Crossing pedestrian", RED)
callout(s, *hero_xy("cutin"), IX + 5.55, 4.95, "Two-wheeler cut-in", AMBER)
callout(s, *hero_xy("workzone"), IX + 1.25, 3.55, "Temporary obstacle", GREY)
callout(s, *hero_xy("cattle"), IX + 6.35, 3.15, "Cattle", GREY)
callout(s, *hero_xy("oncoming"), IX + 5.6, 2.2, "Oncoming, own side", NAVY)

# official "TITLE PAGE" subtitle -> small label above the name
sub = shape_by_name(s, "Subtitle 3")
set_text_keep_format(sub, "TITLE PAGE")
sub.left, sub.top, sub.width, sub.height = Inches(0.45), Inches(1.28), Inches(4.5), Inches(0.4)
for p in sub.text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.size = Pt(16)
        r.font.color.rgb = rgb(MUTED)
sub.text_frame.vertical_anchor = MSO_ANCHOR.TOP
# name + subtitle
text(s, 0.45, 1.68, 5.3, 0.75, "PATHSENSE", size=44, bold=True, color=NAVY)
text(s, 0.47, 2.45, 4.9, 0.8, "Adaptive Path Planning and Collision Avoidance for Autonomous Vehicles on Unstructured Indian Roads",
     size=14, color=TEAL, bold=True, line_spacing=1.05)

# official fields: re-use the template's text box, same bullets, fitted
tb = shape_by_name(s, "TextBox 9")
tb.left, tb.top, tb.width, tb.height = Inches(0.3), Inches(3.55), Inches(5.0), Inches(3.7)
values = {"Problem Statement ID": f" {PS_ID}", "Problem Statement Title": f" {PS_TITLE}", "Theme": " Smart Vehicles",
          "PS Category": " Software", "Team ID": f" {TEAM_ID}" if TEAM_ID else "", "Team Name": f" {TEAM_NAME}"}
for p in list(tb.text_frame.paragraphs):
    full = "".join(r.text for r in p.runs)
    if not full.strip():
        p._p.getparent().remove(p._p)           # leading blank paragraph
        continue
    key = next((k for k in values if full.strip().startswith(k)), None)
    if key is None:
        continue
    if key == "PS Category":                       # "PS Category- Software/Hardware" -> label + value
        for r in p.runs:
            r.text = r.text.replace("Software/Hardware", "").replace("Software/ Hardware", "")
    if key == "Team Name":                         # "Team Name (Registered on portal)" -> "Team Name -"
        for r in p.runs:
            r.text = r.text.replace(" (Registered on portal)", "").replace("(Registered on portal)", "")
        if not p.runs[-1].text.rstrip().endswith(("-", "–")):
            p.runs[-1].text = p.runs[-1].text.rstrip() + " –"
    if values[key]:
        r = p.add_run()
        r.text = values[key]
        r.font.bold = False
        r.font.name = p.runs[0].font.name
        r.font.color.rgb = rgb(INK)
for p in tb.text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.0
    p.space_before = Pt(0)
    p.space_after = Pt(9)
    for r in p.runs:
        r.font.size = Pt(14)
notes(s, "Hero visual: original 3D concept render (presentation/hero/scene.html) - not a prototype screenshot. "
         "Team ID: fill in from the SIH portal.")

# ================================================================== 2. IDEA TITLE / PROPOSED SOLUTION
s = slides[1]
set_text_keep_format(shape_by_name(s, "Title 1"), "PATHSENSE")
remove(shape_by_name(s, "TextBox 8"))
text(s, 0.4, 1.3, 12.5, 0.36, [[("❖ ", {"color": BLUE}), ("Proposed Solution (Describe your Idea/Solution/Prototype)", {"color": NAVY})]],
     size=16, bold=True)
text(s, 0.4, 1.7, 12.5, 0.3, "See every road user, predict where they are going, and keep re-planning a collision-free path — without relying on lane markings.",
     size=12, color=MUTED, bold=False)

L, LW = 0.4, 6.1
blocks = [
    ("Detailed explanation of the proposed solution", [
        "Camera-first perception detects and tracks cars, buses, trucks, auto-rickshaws, two-wheelers, pedestrians and animals",
        "Short-term motion prediction + a bird’s-eye risk map of the scene",
        "Adaptive planner samples candidate paths and keeps the safest collision-free one, every frame"]),
    ("How it addresses the problem", [
        "Plans in drivable free space, so missing lane markings are not a blocker",
        "Reads cut-ins, crossings, wrong-way and sudden turns from motion history",
        "Built for left-hand, mixed Indian traffic"]),
    ("Innovation and uniqueness of the solution", [
        "Threat to the ego path, not mere proximity: traffic on its own side and people on the footpath are watched; anything entering the path is acted on",
        "Explainable: every GO / SLOW / BRAKE carries its reason"]),
]
y = 2.2
for label, pts in blocks:
    pointer(s, L, y, LW, label, size=12.5)
    tb_ = text(s, L + 0.14, y + 0.36, LW - 0.2, 1.2, pts, size=11, color=INK, bullets=True, space=3)
    y += 0.42 + sum(0.2 * (1 + len(p_) // 78) for p_ in pts) + 0.05 * len(pts) + 0.18

# concept render with callouts for the characteristics of unstructured Indian roads
PW = 6.05
PX, PY = 12.93 - PW, 2.2
img = Image.open(ISO)
PH = PW * img.height / img.width
pic = s.shapes.add_picture(ISO, Inches(PX), Inches(PY), Inches(PW), Inches(PH))
pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
_geom = pic._element.spPr.find(qn("a:prstGeom"))
_av = _geom.find(qn("a:avLst"))
if _av is None:
    _av = etree.SubElement(_geom, qn("a:avLst"))
_gd = etree.SubElement(_av, qn("a:gd")); _gd.set("name", "adj"); _gd.set("fmla", "val 4000")
pic.line.color.rgb = rgb(LINE)
pic.line.width = Pt(0.75)


def iso_xy(k):
    ax, ay = A_ISO[k]
    x0, y0, x1, y1 = ISO_BOX
    return PX + (ax - x0) / (x1 - x0) * PW, PY + (ay - y0) / (y1 - y0) * PH


for k, fx, fy, lab, col in [("pedestrian", 0.19, 0.08, "Pedestrians crossing", RED), ("wrongway", 0.15, 0.36, "Wrong-way rider", AMBER),
                            ("cattle", 0.66, 0.07, "Animals on road", GREY), ("auto", 0.45, 0.22, "Auto-rickshaws", NAVY),
                            ("workzone", 0.17, 0.62, "Temporary obstacle", GREY), ("cutin", 0.8, 0.5, "Informal merge / cut-in", AMBER),
                            ("corridor", 0.62, 0.8, "Adaptive safe path", TEAL)]:
    callout(s, *iso_xy(k), PX + fx * PW, PY + fy * PH, lab, col, size=9.5)
# the remaining characteristics as a compact row under the render
cy = PY + PH + 0.14
text(s, PX, cy, 1.4, 0.28, "Also designed for:", size=9.5, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
cx = PX + 1.3
for lab in ["Mixed traffic", "No lane markings", "Two-wheelers", "Sudden turns"]:
    cx += pill(s, cx, cy, lab, NAVY, size=9, h=0.28, fill=TINT, bold=False) + 0.07
notes(s, "Visual: original 3D concept render (presentation/hero/scene.html), left-hand traffic. Not a prototype screenshot.")

# ================================================================== 3. TECHNICAL APPROACH
s = slides[2]
remove(shape_by_name(s, "TextBox 8"))
pointer(s, 0.4, 1.3, 12.5, "Technologies to be used (e.g. programming languages, frameworks, hardware)")
groups = [("Perception", ["Object detection (YOLO family)", "Multi-object tracking", "Monocular depth estimation"], BLUE),
          ("Prediction & planning", ["Short-term motion prediction", "Bird’s-eye risk / cost map", "Kinematic path sampling"], BLUE),
          ("Software", ["Python · PyTorch", "OpenCV", "Web dashboard for review"], BLUE),
          ("Sensors & compute", ["Front camera (prototype)", "Radar / LiDAR (pilot)", "Automotive compute (production)"], BLUE),
          ("Simulation & validation", ["MATLAB / Simulink", "RoadRunner scenarios", "Automated Driving Toolbox"], AMBER)]
gx = 0.4
widths = [2.55, 2.55, 2.1, 2.6, 2.5]
for (title, items, col), gw in zip(groups, widths):
    tt = title + ("  · pilot stage" if col == AMBER else "")
    text(s, gx, 1.72, gw, 0.24, tt, size=10, bold=True, color=AMBER if col == AMBER else MUTED)
    yy = 1.98
    for it in items:
        w = gw - 0.1
        box(s, gx, yy, w, 0.3, fill=TINT if col == BLUE else AMBER_T, line=LINE if col == BLUE else "F3D19C", radius=0.08)
        text(s, gx + 0.1, yy, w - 0.15, 0.3, it, size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        yy += 0.36
    gx += gw + 0.05

pointer(s, 0.4, 3.22, 12.5, "Methodology and process for implementation (Flow Charts/Images/ working prototype)")
flow = [("web", "Road scene", "front camera view of mixed traffic"),
        ("visibility", "Perception", "detect every road user"),
        ("timeline", "Object tracking", "identity and history over time"),
        ("brain", "Motion prediction", "where each agent goes next"),
        ("warning", "Ego-path / risk", "who threatens our path, and how soon"),
        ("route", "Adaptive planning", "sample and score collision-free paths"),
        ("shield", "Safe trajectory", "best path + GO / SLOW / BRAKE"),
        ("speed", "Vehicle control", "steering and speed commands")]
n = len(flow)
fg = 0.16
fw = (12.53 - fg * (n - 1)) / n
fy, fh = 3.64, 1.45
for i, (ic, t1, t2) in enumerate(flow):
    x = 0.4 + i * (fw + fg)
    hl = i in (4, 5, 6)
    ctrl = i == n - 1
    box(s, x, fy, fw, fh, fill=TEAL_T if hl else "FFFFFF", line=TEAL if hl else (GREY if ctrl else LINE), radius=0.1, lw=1.0,
        dash="dash" if ctrl else None)
    icon(s, ic, x + 0.12, fy + 0.13, 0.34, "e" if hl else "b")
    text(s, x + 0.12, fy + 0.5, fw - 0.2, 0.42, t1, size=11, bold=True, color=NAVY, line_spacing=0.9)
    text(s, x + 0.12, fy + 0.95, fw - 0.2, 0.48, t2, size=9, color=MUTED)
    if i < n - 1:
        arrow(s, x + fw + 0.01, fy + fh / 2, x + fw + fg - 0.01, fy + fh / 2, color=BLUE, width=1.5)
# closed re-planning loop (from safe trajectory back to perception)
lx0, lx1, ly = 0.4 + 1 * (fw + fg) + fw / 2, 0.4 + 6 * (fw + fg) + fw / 2, fy + fh + 0.16
line(s, lx1, fy + fh, lx1, ly, color=TEAL, width=1.5)
line(s, lx1, ly, lx0, ly, color=TEAL, width=1.5)
arrow(s, lx0, ly, lx0, fy + fh + 0.01, color=TEAL, width=1.5)
text(s, (lx0 + lx1) / 2 - 2.5, ly + 0.03, 5.0, 0.24, "continuous re-planning as the scene changes", size=9.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.4 + 7 * (fw + fg), fy + fh + 0.05, fw, 0.24, "pilot / production", size=8.5, color=GREY, align=PP_ALIGN.CENTER)

principles = [("route", "Lane-free planning", "Plans in drivable free space, so missing or faded lane markings don’t break it."),
              ("walk", "Threat, not proximity", "Traffic on its own side and footpath pedestrians are watched; agents entering the path trigger action."),
              ("layers", "Modular by design", "Perception, prediction and planning are separate stages — sensors and simulators plug in later.")]
for i, (ic, t1, t2) in enumerate(principles):
    x = 0.4 + i * 4.22
    box(s, x, 5.72, 4.05, 1.04, fill=TINT, radius=0.1)
    icon(s, ic, x + 0.18, 5.88, 0.5, "e" if i == 1 else "b", circle="FFFFFF", pad=0.18)
    text(s, x + 0.84, 5.82, 3.1, 0.28, t1, size=12, bold=True, color=NAVY)
    text(s, x + 0.84, 6.11, 3.1, 0.62, t2, size=10, color=MUTED)
notes(s, "Conceptual pipeline of the proposed system. Vehicle control and the MATLAB/RoadRunner simulation stack are pilot-stage items.")

# ================================================================== 4. FEASIBILITY AND VIABILITY
s = slides[3]
remove(shape_by_name(s, "TextBox 8"))
pointer(s, 0.4, 1.3, 8.0, "Analysis of the feasibility of the idea")
cards = [("FEASIBILITY", "check", BLUE, [
             "Builds on established perception, tracking and motion-planning methods",
             "Camera-first design gives an accessible first deployment",
             "Modular: LiDAR / radar can be added in later stages",
             "Validation grows from recorded roads to simulation and closed-loop tests",
             "Mature automotive simulation and perception ecosystems support scale-up"]),
         ("VIABILITY", "rocket", TEAL, [
             "Designed around real Indian-road traffic complexity",
             "Fits urban roads, village roads, highways and dense mixed traffic",
             "Modular perception / planning supports progressive upgrades",
             "Evolves from research prototype → pilot → production",
             "Can integrate with vehicle systems and mobility platforms"])]
for i, (title, ic, col, pts) in enumerate(cards):
    x = 0.4 + i * 4.05
    box(s, x, 1.7, 3.9, 2.42, fill=TINT if col == BLUE else TEAL_T, radius=0.12)
    icon(s, ic, x + 0.18, 1.82, 0.3, "b" if col == BLUE else "e")
    text(s, x + 0.56, 1.83, 3.0, 0.3, title, size=13, bold=True, color=col)
    text(s, x + 0.2, 2.24, 3.55, 1.85, pts, size=10.5, color=INK, bullets=True, space=3.5, bullet_color=col)

pointer(s, 0.4, 4.3, 3.9, "Potential challenges and risks", size=12)
pointer(s, 4.2, 4.3, 4.2, "Strategies for overcoming these challenges", size=12)
risks = [("Unstructured roads, missing lane markings", "Ego-path estimation and scene-aware planning instead of lane-only planning"),
         ("Irregular traffic behaviour", "Short-term motion prediction and dynamic risk assessment"),
         ("Perception uncertainty", "Confidence-aware decisions, temporal tracking, future multi-sensor fusion"),
         ("Real-world validation complexity", "Scenario library + simulation + progressive field validation")]
ry, rh = 4.68, 0.52
for i, (r_, m_) in enumerate(risks):
    y = ry + i * (rh + 0.04)
    box(s, 0.4, y, 7.95, rh, fill="FFFFFF", line=LINE, radius=0.08)
    icon(s, "warning", 0.52, y + 0.13, 0.26, "a")
    text(s, 0.88, y, 2.95, rh, r_, size=10.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, 3.82, y + rh / 2, 4.1, y + rh / 2, color=BLUE, width=1.25)
    text(s, 4.22, y, 4.05, rh, m_, size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# PROTOTYPE -> PILOT -> PRODUCTION (conceptual; no metrics)
text(s, 8.75, 1.3, 4.2, 0.3, "PROTOTYPE → PILOT → PRODUCTION", size=12, bold=True, color=NAVY)
stages = [("PROTOTYPE", BLUE, TINT2, "Camera-based proof of concept: perception, short-term motion understanding and adaptive collision-free path planning."),
          ("PILOT", TEAL, TEAL_T, "Multi-scenario validation on diverse Indian roads, with additional sensors and closed-loop simulation."),
          ("PRODUCTION", NAVY, "E9ECF2", "Vehicle-grade multi-sensor perception, real-time planning and simulation-backed validation under automotive safety constraints.")]
for i, (t1, col, fill, t2) in enumerate(stages):
    x, y = 8.75 + i * 0.3, 1.75 + i * 1.72
    w = 4.18 - i * 0.3 * 1.0
    box(s, x, y, w, 1.38, fill=fill, radius=0.1)
    text(s, x + 0.2, y + 0.12, w - 0.4, 0.4, t1, size=17, bold=True, color=col)
    text(s, x + 0.2, y + 0.55, w - 0.35, 0.8, t2, size=10.5, color=INK)
    if i < 2:
        c = connector(s, x + 0.35, y + 1.38, x + 0.65, y + 1.72, color=col, width=1.75, kind=MSO_CONNECTOR.CURVE)
notes(s, "Roadmap is conceptual: prototype = current proof of concept; pilot and production are future stages.")

# ================================================================== 5. IMPACT AND BENEFITS
s = slides[4]
remove(shape_by_name(s, "TextBox 8"))
pointer(s, 0.4, 1.3, 12.5, "Potential impact on the target audience")
aud = [("walk", "Road users", "Pedestrians, two-wheeler riders and passengers in mixed traffic"),
       ("car", "OEMs & ADAS teams", "An Indian-road planning layer for driver assistance and autonomy"),
       ("bus", "Fleets & public transport", "Buses, taxis, autos and logistics on unstructured routes"),
       ("science", "Research & testing bodies", "Scenario-based validation of adaptive navigation")]
for i, (ic, t1, t2) in enumerate(aud):
    x = 0.4 + i * 3.16
    box(s, x, 1.7, 3.02, 1.05, fill=TINT, radius=0.1)
    icon(s, ic, x + 0.15, 1.88, 0.5, "b", circle="FFFFFF", pad=0.18)
    text(s, x + 0.78, 1.8, 2.15, 0.3, t1, size=11.5, bold=True, color=NAVY)
    text(s, x + 0.78, 2.1, 2.15, 0.6, t2, size=9.5, color=MUTED)

pointer(s, 0.4, 2.98, 12.5, "Benefits of the solution (social, economic, environmental, etc.)")
FY, FH = 3.4, 3.38
# current challenge
box(s, 0.4, FY, 2.75, FH, fill="F6F7F9", line=LINE, radius=0.12)
text(s, 0.6, FY + 0.15, 2.4, 0.3, "CURRENT CHALLENGE", size=11, bold=True, color=RED)
chal = ["Unstructured roads", "Mixed traffic", "Unpredictable movement", "Limited lane information"]
yy = FY + 0.55
for j, c_ in enumerate(chal):
    text(s, 0.6, yy, 2.4, 0.28, c_, size=12, bold=True, color=NAVY)
    yy += 0.3
    if j < len(chal) - 1:
        text(s, 0.6, yy - 0.04, 2.4, 0.22, "+", size=11, color=GREY)
        yy += 0.2
text(s, 0.6, FY + FH - 0.72, 2.4, 0.6, "India, 2022: 4,61,312 road accidents and 1,68,491 deaths (MoRTH)", size=9, color=MUTED)
arrow(s, 3.2, FY + FH / 2, 3.5, FY + FH / 2, color=BLUE, width=2)
# pathsense
box(s, 3.55, FY, 2.75, FH, fill=TEAL_T, radius=0.12)
text(s, 3.75, FY + 0.15, 2.4, 0.3, "PATHSENSE", size=11, bold=True, color=TEAL)
sol = ["Scene understanding", "Motion prediction", "Risk-aware path planning", "Adaptive re-planning"]
yy = FY + 0.55
for j, c_ in enumerate(sol):
    text(s, 3.75, yy, 2.45, 0.28, c_, size=12, bold=True, color=NAVY)
    yy += 0.3
    if j < len(sol) - 1:
        text(s, 3.75, yy - 0.04, 2.4, 0.22, "+", size=11, color=TEAL)
        yy += 0.2
text(s, 3.75, FY + FH - 0.72, 2.4, 0.6, "Proposed system — benefits are expected outcomes, not measured results", size=9, color=MUTED)
arrow(s, 6.35, FY + FH / 2, 6.65, FY + FH / 2, color=BLUE, width=2)
# expected benefits
box(s, 6.7, FY, 6.2, FH, fill="FFFFFF", line=LINE, radius=0.12)
text(s, 6.9, FY + 0.15, 3.5, 0.3, "EXPECTED BENEFITS", size=11, bold=True, color=BLUE)
ben = [("shield", "Safety", "Social", "Better handling of unpredictable road interactions"),
       ("route", "Mobility", "Economic · Environmental", "Smoother, adaptive navigation where lanes are unreliable"),
       ("layers", "Scalability", "Economic", "Architecture can evolve toward multi-sensor vehicle systems"),
       ("map", "Indian-road relevance", "Social", "Designed around mixed, irregular traffic behaviour"),
       ("science", "Research value", "Research", "A path to scenario-based validation of adaptive navigation")]
for j, (ic, t1, tag, t2) in enumerate(ben):
    y = FY + 0.52 + j * 0.56
    icon(s, ic, 6.92, y + 0.04, 0.3, "b")
    text(s, 7.36, y, 2.0, 0.26, t1, size=11.5, bold=True, color=NAVY)
    text(s, 7.36, y + 0.25, 5.4, 0.26, t2, size=10, color=MUTED)
    tw = 0.24 + 0.058 * len(tag)
    pill(s, 12.72 - tw, y + 0.0, tag, TEAL if tag != "Research" else BLUE, size=8, h=0.22, w=tw, bold=False)
notes(s, "Benefits are expected / potential impact of the proposed system, not measured results. "
         "Accident figures: Ministry of Road Transport & Highways, 'Road Accidents in India 2022'.")

# ================================================================== 6. RESEARCH AND REFERENCES
s = slides[5]
remove(shape_by_name(s, "TextBox 8"))
pointer(s, 0.4, 1.3, 12.5, "Details / Links of the reference and research work")
refs = [("map", "Indian-road perception", ["Varma et al., IDD: A Dataset for Exploring Problems of Autonomous Navigation in Unconstrained Environments, WACV 2019"]),
        ("route", "Path planning", ["Paden et al., A Survey of Motion Planning and Control Techniques for Self-Driving Urban Vehicles, IEEE T-IV 2016",
                                    "González et al., A Review of Motion Planning Techniques for Automated Vehicles, IEEE T-ITS 2016"]),
        ("visibility", "Detection, tracking & depth", ["Redmon et al., You Only Look Once, CVPR 2016",
                                                       "Zhang et al., ByteTrack, ECCV 2022",
                                                       "Yang et al., Depth Anything V2, NeurIPS 2024"]),
        ("brain", "Motion prediction", ["Mozaffari et al., Deep Learning-Based Vehicle Behavior Prediction for Autonomous Driving Applications: A Review, IEEE T-ITS 2022",
                                        "Alahi et al., Social LSTM, CVPR 2016"]),
        ("block", "Collision avoidance", ["Fiorini & Shiller, Motion Planning in Dynamic Environments Using Velocity Obstacles, IJRR 1998",
                                          "Hayward, Near-miss determination through use of a scale of danger (time-to-collision), HRR 1972"]),
        ("shield", "Autonomous-driving safety", ["ISO 26262 — Road vehicles, functional safety",
                                                 "ISO 21448 — Safety of the intended functionality (SOTIF)",
                                                 "Shalev-Shwartz et al., On a Formal Model of Safe and Scalable Self-driving Cars, 2017"]),
        ("location", "Road context", ["OpenStreetMap contributors — road network and points of interest",
                                      "MoRTH, Road Accidents in India 2022"]),
        ("flag", "Problem statement & tools", ["SIH 2026 problem statement SIH26037 — " + PS_TITLE,
                                               "MathWorks RoadRunner & Automated Driving Toolbox documentation"])]
CW, CH, GX, GY = 3.02, 2.45, 0.14, 0.14
for i, (ic, title, items) in enumerate(refs):
    col, row = i % 4, i // 4
    x, y = 0.4 + col * (CW + GX), 1.72 + row * (CH + GY)
    box(s, x, y, CW, CH, fill=TINT if (col + row) % 2 == 0 else "FFFFFF", line=None if (col + row) % 2 == 0 else LINE, radius=0.1)
    icon(s, ic, x + 0.16, y + 0.15, 0.3, "b")
    text(s, x + 0.54, y + 0.17, CW - 0.65, 0.3, title, size=11.5, bold=True, color=NAVY)
    text(s, x + 0.18, y + 0.6, CW - 0.32, CH - 0.7, items, size=10.5, color=INK, bullets=True, space=5)
notes(s, "Concise references by category. Proposed-system claims are conceptual; LiDAR/radar, simulation-backed validation "
         "and vehicle integration are pilot / production roadmap items.")

dst = os.path.join(ROOT, "PathSense_SIH_Presentation.pptx")
prs.save(dst)
print("saved", dst, len(prs.slides), "slides")
